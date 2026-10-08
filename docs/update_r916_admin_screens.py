"""Insert supplied admin screenshots into section 3.9.1 of the author's report."""
import re
from copy import deepcopy
from pathlib import Path
from PIL import Image
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SOURCE = Path('C:/Users/Advice/Downloads/PalmGuard_r916.docx')
ROOT = Path(__file__).resolve().parent
doc = Document(SOURCE)
before_pictures = len(doc.inline_shapes)
anchor = next(p for p in doc.paragraphs if p.text.strip().startswith('3.9.2 '))
body_sample = next(p for p in doc.paragraphs if p.text.lstrip().startswith('จากรูปที่ 28 แสดงหน้าโปรไฟล์'))
caption_sample = next(p for p in doc.paragraphs if p.text.strip().startswith('รูปที่ 28 หน้าโปรไฟล์'))

# Replace figure references within their original XML text nodes, retaining styles,
# hyperlinks, fields, section settings, and the rest of the source document.
for node in doc.element.iter(qn('w:t')):
    if node.text:
        node.text = re.sub(r'(รูปที่\s*)(\d+)(?![\d.])', lambda m: m[1] + str(int(m[2])+4) if int(m[2]) >= 30 else m[0], node.text)

def insert(text='', caption=False):
    element = OxmlElement('w:p')
    anchor._p.addprevious(element)
    from docx.text.paragraph import Paragraph
    p = Paragraph(element, anchor._parent)
    sample = caption_sample if caption else body_sample
    if sample._p.pPr is not None:
        element.append(deepcopy(sample._p.pPr))
    p.style = sample.style
    p.paragraph_format.first_line_indent = Cm(0 if caption else 1.25)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if caption else WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.keep_with_next = False
    p.paragraph_format.space_after = Pt(6)
    if text:
        run = p.add_run(text)
        run.font.name = 'TH Sarabun New'
        run.font.size = Pt(16)
        run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'TH Sarabun New')
    return p

insert('ส่วนติดต่อผู้ดูแลระบบออกแบบสำหรับผู้ใช้งานที่ได้รับสิทธิ์ผู้ดูแลระบบ เพื่อเข้าถึงข้อมูลบัญชีผู้ใช้งาน ผลการตรวจวิเคราะห์ และคำตอบแบบประเมินความพึงพอใจ โดยมีเมนูเชื่อมโยงระหว่างหน้าภาพรวมระบบ หน้าผู้ใช้งาน หน้าผลตรวจวิเคราะห์ และหน้าแบบประเมิน เพื่อให้ตรวจสอบข้อมูลแต่ละส่วนได้สะดวก')

blocks = [
('ad2.png', 'หน้าโปรไฟล์ผู้ดูแลระบบ',
 'หน้าโปรไฟล์ผู้ดูแลระบบแสดงข้อมูลบัญชี ประกอบด้วยชื่อผู้ใช้งาน อีเมล บทบาท วันที่สมัครสมาชิก และจำนวนรายการตรวจวิเคราะห์ พร้อมส่วนข้อมูลส่วนตัวสำหรับตรวจสอบและแก้ไขข้อมูลบัญชี ภายในหน้ามีลิงก์ไปยังประวัติการวินิจฉัยและส่วนจัดการระบบ รวมถึงปุ่มเข้าสู่แบบสอบถามความพึงพอใจและปุ่มออกจากระบบ',
 'แสดงหน้าโปรไฟล์ของผู้ดูแลระบบ ซึ่งช่วยให้ผู้ดูแลตรวจสอบข้อมูลบัญชีของตนเอง และเข้าถึงส่วนจัดการระบบผ่านสิทธิ์ที่ได้รับ โดยการแก้ไขข้อมูลส่วนตัวทำผ่านปุ่มแก้ไขข้อมูลภายในหน้า'),
('ad3.png', 'หน้าจัดการผู้ใช้งาน',
 'หน้าจัดการผู้ใช้งานแสดงรายชื่อบัญชีที่ลงทะเบียนในระบบ โดยแสดงชื่อ อีเมล สิทธิ์การใช้งาน วันที่สมัครสมาชิก และสถานะการตอบแบบประเมินความพึงพอใจ ผู้ดูแลระบบสามารถเลือกปุ่มดูรายละเอียดเพื่อเข้าถึงข้อมูลของผู้ใช้งานและประวัติการตรวจวิเคราะห์ของบัญชีนั้นได้',
 'แสดงหน้าจัดการผู้ใช้งาน ซึ่งช่วยให้ผู้ดูแลระบบตรวจสอบข้อมูลบัญชีและแยกบทบาทระหว่างเกษตรกรกับผู้ดูแลระบบได้ พร้อมติดตามสถานะการตอบแบบประเมินของแต่ละบัญชีจากรายการที่แสดง'),
('ad4.png', 'หน้าจัดการผลการตรวจวิเคราะห์',
 'หน้าจัดการผลการตรวจวิเคราะห์แสดงรายการผลตรวจของผู้ใช้งานภายในระบบ พร้อมเครื่องมือค้นหาจากชื่อ อีเมล หรือรหัสรายการ และตัวกรองตามสถานะและประเภทผลตรวจ แต่ละรายการแสดงภาพขนาดย่อ ชื่อผู้ใช้งาน ผลการจำแนก ค่าความเชื่อมั่นเมื่อวิเคราะห์สำเร็จ และวันเวลาที่ตรวจวิเคราะห์ ผู้ดูแลระบบสามารถดูรายละเอียด แก้ไข และลบรายการตามสิทธิ์ที่ระบบกำหนดได้',
 'แสดงหน้าจัดการผลการตรวจวิเคราะห์ ซึ่งช่วยให้ผู้ดูแลระบบค้นหาและคัดกรองรายการที่ต้องการตรวจสอบได้สะดวก โดยแสดงรายการวิเคราะห์สำเร็จและไม่สำเร็จแยกตามสถานะ รวมทั้งมีปุ่มสำหรับจัดการข้อมูลและเครื่องมือแบ่งหน้าเพื่อดูรายการย้อนหลัง'),
('ad5.png', 'หน้าภาพรวมแบบประเมินความพึงพอใจ',
 'หน้าภาพรวมแบบประเมินความพึงพอใจแสดงจำนวนคำตอบ จำนวนบัญชีที่ตอบแบบประเมิน และสัดส่วนการแนะนำระบบแก่ผู้อื่น พร้อมคะแนนเฉลี่ยและกราฟแจกแจงคะแนนระดับ 1–5 ในด้านความพึงพอใจโดยรวม ความง่ายในการใช้งาน และความคิดเห็นต่อความแม่นยำของระบบ ด้านล่างแสดงรายการคำตอบและความคิดเห็นรายบุคคล โดยผู้ดูแลระบบสามารถเปิดดูคำตอบได้ แต่ไม่มีเครื่องมือแก้ไขคำตอบของผู้ใช้งาน',
 'แสดงหน้าภาพรวมแบบประเมินความพึงพอใจ ซึ่งช่วยให้ผู้ดูแลระบบพิจารณาคะแนนแต่ละด้านร่วมกับความคิดเห็นรายบุคคลเพื่อนำไปปรับปรุงระบบ ทั้งนี้ คะแนนความคิดเห็นต่อความแม่นยำเป็นการรับรู้ของผู้ตอบแบบประเมิน และแยกจากค่า Accuracy ที่ได้จากการทดสอบแบบจำลอง'),
]
for number, (filename, title, introduction, explanation) in enumerate(blocks, 30):
    insert(introduction)
    picture = insert(caption=True)
    picture.paragraph_format.keep_with_next = True
    path = Path('C:/Users/Advice/Downloads') / filename
    width_px, height_px = Image.open(path).size
    width_cm = min(15.5, 19 * width_px / height_px)
    picture.add_run().add_picture(str(path), width=Cm(width_cm))
    insert(f'รูปที่ {number} {title}', caption=True)
    insert(f'จากรูปที่ {number} {explanation}')

# Request a field refresh when the report is opened in Word.
settings = doc.settings.element
update = settings.find(qn('w:updateFields'))
if update is None:
    update = OxmlElement('w:updateFields')
    settings.append(update)
update.set(qn('w:val'), 'true')
output = ROOT / 'PalmGuard_r916_admin_updated.docx'
doc.save(output)
check = Document(output)
assert len(check.inline_shapes) == before_pictures + 4
texts = [p.text.strip() for p in check.paragraphs]
start = next(i for i, text in enumerate(texts) if text.startswith('3.9.1 '))
end = next(i for i, text in enumerate(texts) if text.startswith('3.9.2 '))
for number, (_, title, _, _) in enumerate(blocks, 30):
    assert f'รูปที่ {number} {title}' in texts[start:end]
print('Created and verified updated report with four additional images:', output)
