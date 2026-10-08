"""Repair caption numbering, local references and bookmark-backed lists."""
import re
from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.text.paragraph import Paragraph
from docx.shared import Pt, Cm
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SOURCE = Path('C:/Users/Advice/Downloads/PalmGuard_s.docx')
OUTPUT = Path(__file__).resolve().parent / 'PalmGuard_s_numbering_fixed.docx'
doc = Document(SOURCE)
elements = list(doc.element.body)
def text(el):
    return ''.join(n.text or '' for n in el.iter(qn('w:t')))

def replace(el, pattern, replacement):
    nodes = list(el.iter(qn('w:t')))
    original = ''.join(n.text or '' for n in nodes)
    bounds = []
    pos = 0
    for node in nodes:
        length = len(node.text or '')
        bounds.append((pos, pos+length, node))
        pos += length
    for match in reversed(list(re.finditer(pattern, original))):
        value = replacement(match) if callable(replacement) else replacement
        touched = [(start,end,node) for start,end,node in bounds if start < match.end() and end > match.start()]
        if not touched:
            continue
        start, end, first = touched[0]
        _, _, last = touched[-1]
        prefix = (first.text or '')[:match.start()-start]
        last_start = touched[-1][0]
        suffix = (last.text or '')[match.end()-last_start:]
        if first is last:
            first.text = prefix + value + suffix
        else:
            first.text = prefix + value
            for _,_,node in touched[1:-1]:
                node.text = ''
            last.text = suffix

caption_pattern = re.compile(r'^\s*(รูปที่|ภาพที่|ตารางที่)\s*(\d+)\s+(.+)', re.S)
captions = []
counts = {'figure':0,'table':0}
for index, el in enumerate(elements):
    if index < 200 or el.tag != qn('w:p'):
        continue
    match = caption_pattern.match(text(el))
    if not match:
        continue
    kind = 'table' if match[1] == 'ตารางที่' else 'figure'
    counts[kind] += 1
    title = match[3].strip()
    if index == 717:
        title = 'หน้าประวัติการตรวจวินิจฉัย'
    if index == 981:
        title = 'หน้าต่างครอปภาพสำหรับเลือกบริเวณใบปาล์ม'
    # A separate explanation was accidentally joined to the final table caption.
    if 'คะแนนความคิดเห็นต่อความแม่นยำ' in title:
        title = title.split('คะแนนความคิดเห็นต่อความแม่นยำ')[0].strip()
    captions.append({'index':index,'element':el,'kind':kind,'old':int(match[2]),'new':counts[kind], 'title':title})

def target(index, kind, old):
    special = {(651,'figure',28):655, (755,'table',45):754}
    desired = special.get((index,kind,old))
    candidates = [c for c in captions if c['kind']==kind and (c['index']==desired if desired is not None else c['old']==old)]
    if not candidates:
        return str(old)
    return str(min(candidates,key=lambda c:abs(c['index']-index))['new'])

caption_elements = {c['element'] for c in captions}
for index, el in enumerate(elements):
    if index < 200 or el in caption_elements:
        continue
    replace(el, r'(รูปที่|ภาพที่|ตารางที่)(\s*)(\d+)(?![\d.])',
            lambda m: ('ตารางที่' if m[1]=='ตารางที่' else 'รูปที่') + m[2] + target(index,'table' if m[1]=='ตารางที่' else 'figure',int(m[3])))
    if re.search(r'รูปที่\s*…',text(el)):
        preceding = [c for c in captions if c['kind']=='figure' and c['index']<index]
        following = [c for c in captions if c['kind']=='figure' and c['index']>index]
        chosen = preceding[-1] if text(el).strip().startswith('คำอธิบาย') else following[0]
        replace(el,r'รูปที่\s*…',f"รูปที่ {chosen['new']}")
    if index == 752:
        replacement = next(c['new'] for c in captions if c['index']==754)
        replace(el,r'แสดงตารางที่\s*$',f'แสดงดังตารางที่ {replacement}')

for c in captions:
    label = 'ตารางที่' if c['kind']=='table' else 'รูปที่'
    el = c['element']
    if any('SEQ' in (n.text or '') for n in el.iter(qn('w:instrText'))):
        # Freeze old isolated sequence fields so Word cannot reset the new numbering.
        for node in list(el.iter(qn('w:instrText'))) + list(el.iter(qn('w:fldChar'))):
            node.getparent().remove(node)
    match = caption_pattern.match(text(el))
    old_title = match[3].strip()
    replace(el,r'^\s*(รูปที่|ภาพที่|ตารางที่)\s*\d+',f"{label} {c['new']}")
    if c['index'] in [717,981]:
        replace(el,re.escape(old_title),c['title'])
    if c['index']==1114:
        phrase = 'คะแนนความคิดเห็นต่อความแม่นยำ'
        full = text(el)
        explanation = full[full.index(phrase):]
        replace(el,re.escape(explanation),'')
        p = OxmlElement('w:p')
        if el.find(qn('w:pPr')) is not None:
            p.append(deepcopy(el.find(qn('w:pPr'))))
        el.addnext(p)
        run = OxmlElement('w:r'); t=OxmlElement('w:t'); t.text=explanation
        run.append(t); p.append(run)

# Repair appendix menu numbering where the crop screen was inserted as item 2.
in_appendix = False
item = 0
for index,el in enumerate(elements):
    if text(el).strip()=='ภาคผนวก (ก)':
        in_appendix=True
    if text(el).strip()=='ภาคผนวก (ข)':
        in_appendix=False
    if in_appendix and re.match(r'^\s*\d+\.\s*หน้า',text(el)):
        item+=1
        replace(el,r'^\s*\d+\.',f'{item}.')

# Use live PAGEREF fields pointing to actual captions instead of stale TOC entries.
bookmark_ids = [int(n.get(qn('w:id'))) for n in doc.element.iter(qn('w:bookmarkStart'))]
next_id = max(bookmark_ids,default=0)+1
for c in captions:
    c['bookmark'] = f"PG_{c['kind']}_{c['new']}"
    start=OxmlElement('w:bookmarkStart'); start.set(qn('w:id'),str(next_id)); start.set(qn('w:name'),c['bookmark'])
    end=OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'),str(next_id)); next_id+=1
    c['element'].insert(1 if c['element'].find(qn('w:pPr')) is not None else 0,start)
    c['element'].append(end)

figure_anchor=elements[170]  # Existing list-of-tables heading.
table_anchor=elements[186]   # First paragraph following old list entries.
for index,el in enumerate(elements[:186]):
    value=text(el).strip()
    if caption_pattern.match(value) or (index>113 and value=='สารบัญรูปภาพ (ต่อ)'):
        el.getparent().remove(el)

def list_entry(anchor,c):
    el=OxmlElement('w:p'); anchor.addprevious(el)
    p=Paragraph(el,doc._body)
    p.paragraph_format.space_after=Pt(3)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(15),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
    label='รูปที่' if c['kind']=='figure' else 'ตารางที่'
    r=p.add_run(f"{label} {c['new']} {c['title']}\t")
    r.font.name='TH Sarabun New'; r.font.size=Pt(16)
    field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),f" PAGEREF {c['bookmark']} \\h ")
    r=OxmlElement('w:r'); t=OxmlElement('w:t'); t.text='—'; r.append(t); field.append(r)
    el.append(field)
for c in captions:
    list_entry(figure_anchor if c['kind']=='figure' else table_anchor,c)

update=doc.settings.element.find(qn('w:updateFields'))
if update is None:
    update=OxmlElement('w:updateFields'); doc.settings.element.append(update)
update.set(qn('w:val'),'true')
doc.save(OUTPUT)
check=Document(OUTPUT)
assert len(check.inline_shapes)==len(Document(SOURCE).inline_shapes)
assert len(check.tables)==len(Document(SOURCE).tables)
assert 'Error! Bookmark not defined.' not in '\n'.join(p.text for p in check.paragraphs)
assert not any('รูปที่ …' in p.text for p in check.paragraphs)
audit=['# Caption numbering repair',f"Figures: {counts['figure']}; tables: {counts['table']}"]
audit += [f"{c['kind']} {c['old']} -> {c['new']}: {c['title']}" for c in captions]
(OUTPUT.parent/'PalmGuard_s_numbering_changes.txt').write_text('\n'.join(audit),encoding='utf-8')
print(f"Saved {OUTPUT}; {counts}; image and table objects preserved")
