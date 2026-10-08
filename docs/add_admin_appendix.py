"""Append annotated admin screens and numbered control descriptions."""
from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.text.paragraph import Paragraph
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

DOWNLOADS = Path('C:/Users/Advice/Downloads')
ROOT = Path(__file__).resolve().parent
doc = Document(DOWNLOADS / 'PalmGuard_r916_admin_updated.docx')
initial_images = len(doc.inline_shapes)
anchor = next(p for p in doc.paragraphs if p.text.strip() == 'ภาคผนวก (ข)')
sample = next(p for p in doc.paragraphs if p.text.startswith('หมายเลข 1 แสดงชื่อบัญชี'))

def add(text='', centered=False):
    element = OxmlElement('w:p')
    anchor._p.addprevious(element)
    paragraph = Paragraph(element, anchor._parent)
    if sample._p.pPr is not None:
        element.append(deepcopy(sample._p.pPr))
    paragraph.style = 'Normal'
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_with_next = False
    paragraph.paragraph_format.page_break_before = False
    if text:
        run = paragraph.add_run(text)
        run.font.name = 'TH Sarabun New'
        run.font.size = Pt(16)
    return paragraph

screens = [
    ('1.png', (403,214,1447,930), '8. หน้าโปรไฟล์ผู้ดูแลระบบ', 'หน้าโปรไฟล์ผู้ดูแลระบบ', [
        'แสดงชื่อ–นามสกุลของผู้ดูแลระบบในส่วนข้อมูลส่วนตัว',
        'แสดงอีเมลที่ใช้สำหรับบัญชีผู้ดูแลระบบ',
        'แสดงหมายเลขโทรศัพท์ของบัญชี หากยังไม่ได้บันทึกจะแสดงข้อความไม่ระบุ',
        'ปุ่มแก้ไขข้อมูล ใช้เปิดแบบฟอร์มเพื่อปรับปรุงข้อมูลส่วนตัวของบัญชี',
        'ปุ่มประวัติการวินิจฉัย ใช้เข้าสู่หน้ารายการตรวจวิเคราะห์ย้อนหลังของบัญชี',
        'ปุ่มจัดการระบบ ใช้เข้าสู่ส่วนผู้ดูแลระบบเพื่อดูและจัดการข้อมูลตามสิทธิ์',
        'ปุ่มทำแบบสอบถาม ใช้เข้าสู่แบบประเมินความพึงพอใจในการใช้งานระบบ',
        'ปุ่มออกจากระบบ ใช้สิ้นสุดการเข้าสู่ระบบของบัญชีที่กำลังใช้งาน',
    ]),
    ('3b4ee5ad-3779-4d34-8d60-2981da5e89d5.jpg', (806,134,1114,947), '9. หน้าจัดการผลการตรวจวิเคราะห์', 'หน้าจัดการผลการตรวจวิเคราะห์สำหรับผู้ดูแลระบบ', [
        'เมนูผลตรวจวิเคราะห์ ใช้เข้าสู่หน้าจัดการรายการตรวจวิเคราะห์ของผู้ใช้งาน',
        'ช่องค้นหา ใช้ระบุชื่อผู้ใช้งาน อีเมล หรือรหัสรายการที่ต้องการค้นหา',
        'ตัวกรองสถานะ ใช้เลือกแสดงรายการตามสถานะของการตรวจวิเคราะห์',
        'ตัวกรองประเภทผลตรวจ ใช้เลือกแสดงรายการตามประเภทผลการจำแนก',
        'ปุ่มล้างตัวกรอง ใช้ยกเลิกคำค้นหาและตัวกรองที่เลือกไว้เพื่อแสดงรายการทั้งหมด',
        'ปุ่มดูรายละเอียด ใช้เปิดภาพและรายละเอียดผลการตรวจวิเคราะห์ของรายการที่เลือก',
        'ปุ่มแก้ไข ใช้เปิดหน้าต่างปรับปรุงข้อมูลผลตรวจตามสิทธิ์ของผู้ดูแลระบบ',
        'ปุ่มลบ ใช้ลบรายการผลตรวจที่เลือก โดยระบบแสดงหน้าต่างยืนยันก่อนลบ',
    ]),
    ('6e8be786-6651-4ba9-9c6a-8594414ad6d1.jpg', (465,145,1455,938), '10. หน้าจัดการผู้ใช้งาน', 'หน้าจัดการผู้ใช้งานสำหรับผู้ดูแลระบบ', [
        'ปุ่มดูรายละเอียด ใช้เปิดข้อมูลของผู้ใช้งานในแถวที่เลือก เพื่อดูข้อมูลบัญชี ประวัติการตรวจวิเคราะห์ และคำตอบแบบประเมินของผู้ใช้งานนั้น',
    ]),
]
for number, (filename, box, heading, title, items) in enumerate(screens,55):
    p = add(heading)
    p.runs[0].bold = True
    p.paragraph_format.page_break_before = True
    p.paragraph_format.keep_with_next = True
    p = add(f'หน้าจอนี้แสดง{title} โดยองค์ประกอบและปุ่มที่กำกับหมายเลขมีหน้าที่ดังรูปที่ {number}')
    p.paragraph_format.keep_with_next = True
    path = DOWNLOADS / filename
    source_w, source_h = Image.open(path).size
    left,top,right,bottom = box
    w,h = right-left,bottom-top
    display_w = min(15.5, 16.7*w/h)
    p = add(centered=True)
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(path), width=Cm(display_w))
    shape.height = Cm(display_w*h/w)
    # Native Word picture cropping removes empty margins without modifying the original image.
    blip_fill = shape._inline.graphic.graphicData.pic.blipFill
    crop = OxmlElement('a:srcRect')
    for key,value in [('l',left/source_w),('t',top/source_h),('r',(source_w-right)/source_w),('b',(source_h-bottom)/source_h)]:
        crop.set(key,str(round(value*100000)))
    blip_fill.insert(1,crop)
    add(f'รูปที่ {number} {title}',centered=True)
    p = add(f'คำอธิบายรูปที่ {number}')
    p.paragraph_format.keep_with_next = True
    for index,text in enumerate(items,1):
        add(f'หมายเลข {index} {text}')

output = ROOT / 'PalmGuard_r916_appendix_admin_updated.docx'
doc.save(output)
check = Document(output)
assert len(check.inline_shapes) == initial_images+3
texts = [p.text for p in check.paragraphs]
b_index = texts.index('ภาคผนวก (ข)')
assert all(f'รูปที่ {n} {title}' in texts[:b_index] for n,(_,_,_,title,_) in enumerate(screens,55))
assert sum(p.text.startswith('หมายเลข ') for p in check.paragraphs) == sum(p.text.startswith('หมายเลข ') for p in Document(DOWNLOADS / 'PalmGuard_r916_admin_updated.docx').paragraphs)+17
print('Created and verified:',output)
