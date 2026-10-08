"""Create the editable project-report section requested by the author."""
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin = section.bottom_margin = Cm(2.54)
section.left_margin, section.right_margin = Cm(3), Cm(2.54)
normal = doc.styles['Normal']
normal.font.name = 'TH Sarabun New'
normal.font.size = Pt(16)
normal.element.rPr.rFonts.set(qn('w:eastAsia'), 'TH Sarabun New')
normal.element.rPr.rFonts.set(qn('w:cs'), 'TH Sarabun New')
normal.paragraph_format.line_spacing = 1
normal.paragraph_format.space_after = Pt(6)

def paragraph(text, indent=True):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if indent:
        p.paragraph_format.first_line_indent = Cm(1.25)
    return p

heading = paragraph('4.3.2 การเตรียมและแบ่งชุดข้อมูล', False)
heading.runs[0].bold = True
heading.paragraph_format.keep_with_next = True

paragraph('การเตรียมข้อมูลภาพสำหรับการพัฒนาแบบจำลองจำแนกประเภทใบปาล์มน้ำมัน ประกอบด้วยข้อมูลภาพชุดเดิมและข้อมูลภาพชุดใหม่ โดยแบ่งภาพออกเป็น 3 ประเภท ได้แก่ ใบจุดสีน้ำตาล (Brown Spots) ใบปกติ (Healthy) และเพลี้ยหอย (White Scale) ก่อนนำข้อมูลไปใช้ในการฝึกและทดสอบ ได้ตรวจสอบความสมบูรณ์ของไฟล์ภาพและคัดแยกภาพซ้ำออกจำนวน 275 ภาพ โดยไม่พบไฟล์ภาพเสียหาย หลังจากตรวจสอบแล้วเหลือข้อมูลภาพทั้งหมด 5,057 ภาพ แบ่งเป็นข้อมูลภาพชุดเดิมจำนวน 2,628 ภาพ และข้อมูลภาพชุดใหม่จำนวน 2,429 ภาพ')
paragraph('ข้อมูลภาพดังกล่าวแบ่งออกเป็นชุดฝึกสอน (Training Set) ประมาณร้อยละ 70 ชุดตรวจสอบระหว่างการฝึก (Validation Set) ประมาณร้อยละ 15 และชุดทดสอบ (Testing Set) ประมาณร้อยละ 15 โดยชุดฝึกสอนใช้สำหรับให้แบบจำลองเรียนรู้ลักษณะของภาพแต่ละประเภท ชุดตรวจสอบใช้สำหรับติดตามผลระหว่างการฝึก และชุดทดสอบใช้สำหรับประเมินผลการจำแนกประเภทภายหลังการฝึก มีจำนวนภาพในแต่ละส่วนดังตารางที่ 10')
caption = paragraph('ตารางที่ 10 จำนวนภาพที่ใช้ในการฝึกสอน ตรวจสอบ และทดสอบแบบจำลองจำแนกประเภทใบปาล์มน้ำมัน', False)
caption.paragraph_format.keep_with_next = True
rows = [
    ['ชุดข้อมูล', 'ส่วนแบ่งข้อมูล', 'ใบจุดสีน้ำตาล\n(ภาพ)', 'ใบปกติ\n(ภาพ)', 'เพลี้ยหอย\n(ภาพ)', 'รวม\n(ภาพ)'],
    ['ชุดเดิม', 'ชุดฝึกสอน', '329', '841', '668', '1,838'],
    ['ชุดเดิม', 'ชุดตรวจสอบ', '70', '180', '143', '393'],
    ['ชุดเดิม', 'ชุดทดสอบ', '71', '181', '145', '397'],
    ['ชุดใหม่', 'ชุดฝึกสอน', '665', '453', '582', '1,700'],
    ['ชุดใหม่', 'ชุดตรวจสอบ', '142', '98', '124', '364'],
    ['ชุดใหม่', 'ชุดทดสอบ', '143', '97', '125', '365'],
    ['รวม', 'ทุกส่วนแบ่ง', '1,420', '1,850', '1,787', '5,057'],
]
table = doc.add_table(rows=0, cols=6)
table.style = 'Table Grid'
table.autofit = False
for index, values in enumerate(rows):
    row = table.add_row()
    no_split = OxmlElement('w:cantSplit')
    row._tr.get_or_add_trPr().append(no_split)
    if index == 0:
        repeat = OxmlElement('w:tblHeader')
        row._tr.get_or_add_trPr().append(repeat)
    for cell, value, width in zip(row.cells, values, [2, 2.9, 3.1, 2.4, 2.5, 2.4]):
        cell.width = Cm(width)
        cell.text = value
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            for run in p.runs:
                run.font.size = Pt(14)
                run.bold = index in [0, len(rows)-1]
        if index in [0, len(rows)-1]:
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), 'D9D9D9' if index == 0 else 'EEEEEE')
            cell._tc.get_or_add_tcPr().append(shade)

paragraph('จากตารางที่ 10 พบว่าข้อมูลภาพที่ใช้ในการฝึกสอนมีจำนวนรวม 3,538 ภาพ ชุดตรวจสอบระหว่างการฝึกมีจำนวน 757 ภาพ และชุดทดสอบมีจำนวน 762 ภาพ โดยชุดทดสอบแบ่งเป็นข้อมูลภาพชุดเดิมจำนวน 397 ภาพ และข้อมูลภาพชุดใหม่จำนวน 365 ภาพ เพื่อประเมินความสามารถของแบบจำลองในการจำแนกภาพจากข้อมูลทั้งสองชุดแยกกัน จำนวนภาพแต่ละประเภทในชุดทดสอบแสดงดังรูปที่ 45')
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.keep_with_next = True
p.add_run().add_picture(str(ROOT / 'report-figures/class-counts-v1.3.png'), width=Cm(15.3))
p = doc.add_paragraph('รูปที่ 45 จำนวนภาพแต่ละประเภทในชุดทดสอบข้อมูลเดิมและข้อมูลใหม่')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
paragraph('จากรูปที่ 45 ชุดทดสอบข้อมูลเดิมประกอบด้วยภาพใบจุดสีน้ำตาลจำนวน 71 ภาพ ภาพใบปกติจำนวน 181 ภาพ และภาพเพลี้ยหอยจำนวน 145 ภาพ ส่วนชุดทดสอบข้อมูลใหม่ประกอบด้วยภาพใบจุดสีน้ำตาลจำนวน 143 ภาพ ภาพใบปกติจำนวน 97 ภาพ และภาพเพลี้ยหอยจำนวน 125 ภาพ โดยใช้จำนวนภาพดังกล่าวเป็นข้อมูลประกอบการประเมินผลการจำแนกประเภทของแบบจำลองในหัวข้อถัดไป')
paragraph('เนื่องจากจำนวนภาพในชุดฝึกสอนของแต่ละประเภทมีความแตกต่างกัน จึงกำหนดค่าน้ำหนักของแต่ละประเภท (Class Weight) เพื่อปรับน้ำหนักในการคำนวณค่าความสูญเสียระหว่างการฝึก โดยประเภทใบจุดสีน้ำตาลมีค่าน้ำหนักเท่ากับ 1.1865 ประเภทใบปกติเท่ากับ 0.9114 และประเภทเพลี้ยหอยเท่ากับ 0.9435 ทั้งนี้ การแบ่งชุดข้อมูลเดิมเป็นการจัดแบ่งย้อนหลังตามข้อมูลประกอบการทดลอง จึงพิจารณาผลการทดสอบเป็นผลประเมินเบื้องต้น')

output = ROOT / 'PalmGuard_section_4.3.2.docx'
doc.save(output)
check = Document(output)
assert len(check.tables) == 1 and len(check.tables[0].rows) == 8
assert len(check.inline_shapes) == 1
assert '4.3.2' in check.paragraphs[0].text
print(f'Created and verified: {output}')
